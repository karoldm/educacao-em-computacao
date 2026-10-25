# Gera slides/aula_pdi.pptx com simulações animadas (GIFs) feitas a partir da foto real em assets/folhas.png.
# Requer: pip install pillow numpy python-pptx. Uso: python3 tools/build_slides.py
import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BG, AC, AMB, INK, MU, CELL = (243, 245, 251), (79, 70, 229), (245, 158, 11), (29, 36, 51), (90, 100, 120), (226, 230, 245)
def F(s, b=False):
    try: return ImageFont.truetype('DejaVuSans-Bold.ttf' if b else 'DejaVuSans.ttf', s)
    except OSError: return ImageFont.load_default()
def FM(s):
    try: return ImageFont.truetype('DejaVuSansMono-Bold.ttf', s)
    except OSError: return F(s, True)
tc = lambda v: (0, 0, 0) if v > 127 else (255, 255, 255)
photo = Image.open('assets/folhas.png').convert('L')
S = np.array(photo.resize((96, 72), Image.LANCZOS)).astype(int)  # mesma resolução da plataforma
small = Image.fromarray(S.astype('uint8'))
def rgb(a): return Image.fromarray(np.clip(a, 0, 255).astype('uint8')).convert('RGB')
def save(frames, durs, name):
    fr = [f.convert('P', palette=Image.ADAPTIVE, colors=48, dither=Image.Dither.NONE) for f in frames]
    fr[0].save('slides/gifs/' + name, save_all=True, append_images=fr[1:], duration=durs, loop=0, optimize=True)
    return 'slides/gifs/' + name
def canvas(w, h): im = Image.new('RGB', (w, h), BG); return im, ImageDraw.Draw(im)
def grid(d, M, x0, y0, cell, font, blank_from=None, hl=None):
    for i, v in enumerate(M.flatten()):
        r, c = divmod(i, M.shape[1]); x, y = x0 + c * cell, y0 + r * cell
        if blank_from is not None and i >= blank_from:
            d.rectangle([x, y, x + cell, y + cell], fill=CELL, outline='white'); d.text((x + cell / 2, y + cell / 2), '?', font=font, fill=MU, anchor='mm')
        else:
            d.rectangle([x, y, x + cell, y + cell], fill=(int(v),) * 3, outline='white'); d.text((x + cell / 2, y + cell / 2), str(int(v)), font=font, fill=tc(v), anchor='mm')
    if hl is not None:
        r, c = divmod(hl, M.shape[1]); d.rectangle([x0 + c * cell, y0 + r * cell, x0 + (c + 1) * cell, y0 + (r + 1) * cell], outline=AMB, width=4)
def caption(im, t, h=44):
    d = ImageDraw.Draw(im); d.rectangle([0, im.height - h, im.width, im.height], fill=INK); d.text((16, im.height - h / 2), t, font=F(20, True), fill='white', anchor='lm')
# trecho 8x8 com bastante branco (mostra a saturação) e trecho 12x9 com contornos fortes (filtro de contorno)
px0, py0 = max(((x, y) for x in range(89) for y in range(65)), key=lambda p: min((S[p[1]:p[1]+8, p[0]:p[0]+8] >= 216).sum(), 8) * 10 + S[p[1]:p[1]+8, p[0]:p[0]+8].std())
P8 = S[py0:py0 + 8, px0:px0 + 8]
ex0, ey0 = max(((x, y) for x in range(85) for y in range(63)), key=lambda p: np.abs(np.diff(S[p[1]:p[1]+9, p[0]:p[0]+12], axis=1)).sum())
P12 = S[ey0:ey0 + 9, ex0:ex0 + 12]

def g_zoom():
    W, H = 480, 360; orig = photo.resize((W, H), Image.LANCZOS).convert('RGB'); pix = small.resize((W, H), Image.NEAREST).convert('RGB'); fr, du = [], []
    f = orig.copy(); caption(f, 'Uma foto digital...'); fr.append(f); du.append(1200)
    for a in np.linspace(.15, 1, 7): f = Image.blend(orig, pix, a); caption(f, '...é feita de quadradinhos'); fr.append(f); du.append(110)
    cx, cy = 46, 38
    for u in np.linspace(0, 1, 22):
        w = 96 * (12 / 96) ** u; h = w * .75; x0 = min(max(cx - w / 2, 0), 96 - w); y0 = min(max(cy - h / 2, 0), 72 - h)
        f = small.resize((W, H), Image.NEAREST, box=(x0, y0, x0 + w, y0 + h)).convert('RGB')
        if u > .8:
            d = ImageDraw.Draw(f); s = W / w
            for k in range(int(w) + 2): d.line([(k * s, 0), (k * s, H)], fill='white', width=1); d.line([(0, k * s * 1), (W, k * s)], fill='white', width=1)
        caption(f, 'Ampliando: cada quadrado é um pixel'); fr.append(f); du.append(100)
    du[-1] = 2000
    return save(fr, du, 'zoom.gif')
def g_matrix():
    fr, du = [], []
    for i in range(64):
        im, d = canvas(800, 500); r, c = divmod(i, 8); v = int(P8[r, c])
        d.text((24, 14), 'Um trecho 8×8 da foto', font=F(20, True), fill=AC)
        grid(d, P8, 24, 50, 52, F(15, True), hl=i)
        d.text((480, 60), f'Pixel: linha {r+1}, coluna {c+1}', font=F(20), fill=INK)
        d.rectangle([480, 100, 600, 220], fill=(v,) * 3, outline=INK, width=2); d.text((690, 160), str(v), font=FM(64), fill=INK, anchor='mm')
        d.text((480, 250), 'número = intensidade do cinza', font=F(18), fill=MU)
        for x in range(300): g = int(x / 299 * 255); d.line([(480 + x, 320), (480 + x, 350)], fill=(g,) * 3)
        mx = 480 + v / 255 * 299; d.polygon([(mx, 356), (mx - 9, 374), (mx + 9, 374)], fill=AMB)
        d.text((480, 392), '0 = preto', font=F(18, True), fill=INK); d.text((780, 392), '255 = branco', font=F(18, True), fill=INK, anchor='ra')
        fr.append(im); du.append(110)
    du[-1] = 1500; return save(fr, du, 'matriz.gif')
def g_gray():
    fr, du = [], []
    for v in list(range(0, 256, 5)) + [255]:
        im, d = canvas(800, 480); d.text((400, 30), 'A escala de cinza', font=F(26, True), fill=AC, anchor='mt')
        d.rectangle([310, 90, 490, 270], fill=(v,) * 3, outline=INK, width=3); d.text((400, 300), str(v), font=FM(54), fill=INK, anchor='mt'); d.text((400, 365), format(v, '08b') + '  (8 bits)', font=FM(22), fill=MU, anchor='mt')
        for x in range(720): g = int(x / 719 * 255); d.line([(40 + x, 410), (40 + x, 450)], fill=(g,) * 3)
        mx = 40 + v / 255 * 719; d.polygon([(mx, 405), (mx - 9, 388), (mx + 9, 388)], fill=AMB)
        d.text((40, 456), '0', font=F(16, True), fill=INK); d.text((760, 456), '255  →  256 níveis possíveis', font=F(16, True), fill=INK, anchor='ra')
        fr.append(im); du.append(70)
    du[-1] = 1500; return save(fr, du, 'cinza.gif')
def g_recipe():
    fr, du = [], []
    for i in range(64):
        im, d = canvas(860, 540); r, c = divmod(i, 8); v = int(P8[r, c]); o = min(255, v + 40)
        d.text((430, 18), 'Regra (receita):  p = p + 40   e se passar de 255, fica 255', font=F(22, True), fill=AC, anchor='mt')
        d.text((30 + 184, 60), 'Entrada', font=F(18, True), fill=INK, anchor='mt'); d.text((462 + 184, 60), 'Saída', font=F(18, True), fill=INK, anchor='mt')
        grid(d, P8, 30, 90, 46, F(13, True), hl=i); out = np.minimum(255, P8 + 40); grid(d, out, 462, 90, 46, F(13, True), blank_from=i + 1, hl=i)
        d.text((410, 270), '→', font=F(40, True), fill=AC, anchor='mm')
        t = f'{v} + 40 = {v+40}' + (f'  →  passou de 255, fica 255!' if v + 40 > 255 else f'  →  {o}')
        d.text((430, 505), t, font=FM(26), fill=INK, anchor='mm'); fr.append(im); du.append(100)
    du[-1] = 1500; return save(fr, du, 'receita.gif')
def g_rules():
    W, H = 520, 400; big = np.array(photo.resize((W, 360), Image.LANCZOS)).astype(int)
    st = [('Original', 'a foto sem filtro', big), ('Clarear em 40', 'p = p + 40  (máx. 255)', np.minimum(255, big + 40)), ('Escurecer em 40', 'p = p − 40  (mín. 0)', np.maximum(0, big - 40)), ('Inverter', 'p = 255 − p', 255 - big)]
    def frame(a, b, t):
        im = Image.new('RGB', (W, H), INK); im.paste(Image.blend(rgb(a[2]), rgb(b[2]), t), (0, 0)); d = ImageDraw.Draw(im); n = b if t > .5 else a
        d.text((16, 380), n[0], font=F(22, True), fill='white', anchor='lm'); d.text((W - 16, 380), n[1], font=FM(18), fill=(200, 200, 255), anchor='rm'); return im
    fr, du = [], []
    for k in range(4):
        a, b = st[k], st[(k + 1) % 4]; fr.append(frame(a, a, 0)); du.append(1400)
        for t in np.linspace(.2, 1, 5): fr.append(frame(a, b, t)); du.append(70)
    return save(fr, du, 'regras.gif')
def edges(a): gx = np.abs(np.roll(a, -1, 1) - np.roll(a, 1, 1)); gy = np.abs(np.roll(a, -1, 0) - np.roll(a, 1, 0)); return np.minimum(255, gx + gy)
def wipe_panel(a, b, w, h): # a: antes, b: depois (arrays HxW já no tamanho)
    return lambda t: (lambda im: (im.paste(rgb(b).crop((int(w * t), 0, w, h)), (int(w * t), 0)), ImageDraw.Draw(im).line([(int(w * t), 0), (int(w * t), h)], fill=AMB, width=3), im)[2])(rgb(a))
def g_uses():
    pw, ph = 280, 210; g = np.array(photo.resize((pw, ph), Image.LANCZOS)).astype(float); rng = np.random.default_rng(3)
    noisy = np.clip(g + rng.normal(0, 28, g.shape), 0, 255); blur = np.array(Image.fromarray(noisy.astype('uint8')).filter(ImageFilter.GaussianBlur(1.6))).astype(float)
    pans = [('Fotografia: foto escura → clarear', g * .42, 255 * (g / 255) ** .55 * .9 + 0), ('Ruído: granulado → suavizar', noisy, blur), ('Contornos: foto → realçar bordas', g, np.minimum(255, 2 * edges(g)))]
    W = [wipe_panel(a, b, pw, ph) for _, a, b in pans]; fr, du = [], []
    for t in list(np.linspace(0, 1, 13)) + list(np.linspace(1, 0, 13))[1:]:
        im, d = canvas(900, 310)
        for k, (n, _, _) in enumerate(pans): x = 10 + k * 295; d.text((x, 12), n, font=F(15, True), fill=AC); im.paste(W[k](t), (x, 50))
        d.text((450, 275), 'antes ◄ linha amarela ► depois', font=F(16), fill=MU, anchor='mt'); fr.append(im); du.append(110 if 0 < t < 1 else 700)
    return save(fr, du, 'usos.gif')
def g_edge_scan():
    fr, du = [], []; c = 34; out = np.zeros((7, 10), int)
    for y in range(1, 8):
        for x in range(1, 11): out[y - 1, x - 1] = min(255, abs(P12[y, x + 1] - P12[y, x - 1]) + abs(P12[y + 1, x] - P12[y - 1, x]))
    n = 0
    for y in range(1, 8):
        for x in range(1, 11):
            im, d = canvas(800, 420); d.text((20, 12), 'Entrada (12×9)', font=F(18, True), fill=INK); d.text((440, 12), 'Saída: contornos', font=F(18, True), fill=INK)
            grid(d, P12, 20, 44, c, F(10, True)); d.rectangle([20 + (x - 1) * c, 44 + (y - 1) * c, 20 + (x + 2) * c, 44 + (y + 2) * c], outline=AMB, width=4)
            flat = np.full(70, 0); flat[:n + 1] = 1; M = out.copy(); grid(d, M, 440, 44 + c, c, F(10, True), blank_from=n + 1); d.rectangle([440 + (n % 10) * c, 44 + c + (n // 10) * c, 440 + (n % 10 + 1) * c, 44 + c + (n // 10 + 1) * c], outline=AMB, width=4)
            a, b = abs(P12[y, x + 1] - P12[y, x - 1]), abs(P12[y + 1, x] - P12[y - 1, x])
            d.text((400, 390), f'|direita − esquerda| + |baixo − cima| = {a} + {b} = {out[y-1,x-1]}', font=FM(20), fill=INK, anchor='mm'); fr.append(im); du.append(110); n += 1
    du[-1] = 1800; return save(fr, du, 'contorno_passo.gif')
def g_edge_wipe():
    w, h = 440, 330; g = np.array(photo.resize((w, h), Image.LANCZOS)).astype(float); W = wipe_panel(g, np.minimum(255, 2 * edges(g)), w, h); fr, du = [], []
    for t in list(np.linspace(0, 1, 14)) + list(np.linspace(1, 0, 14))[1:]: fr.append(W(t)); du.append(110 if 0 < t < 1 else 800)
    return save(fr, du, 'contorno_foto.gif')

def g_blur_scan():
    fr, du = [], []; c = 34; out = np.zeros((7, 10), int)
    for y in range(1, 8):
        for x in range(1, 11): out[y - 1, x - 1] = int(round(P12[y - 1:y + 2, x - 1:x + 2].sum() / 9))
    n = 0
    for y in range(1, 8):
        for x in range(1, 11):
            im, d = canvas(800, 440); d.text((20, 12), 'Entrada (12×9)', font=F(18, True), fill=INK); d.text((440, 12), 'Saída: suavizada', font=F(18, True), fill=INK)
            grid(d, P12, 20, 44, c, F(10, True)); d.rectangle([20 + (x - 1) * c, 44 + (y - 1) * c, 20 + (x + 2) * c, 44 + (y + 2) * c], outline=AMB, width=4)
            grid(d, out, 440, 44 + c, c, F(10, True), blank_from=n + 1); d.rectangle([440 + (n % 10) * c, 44 + c + (n // 10) * c, 440 + (n % 10 + 1) * c, 44 + c + (n // 10 + 1) * c], outline=AMB, width=4)
            v = P12[y - 1:y + 2, x - 1:x + 2].flatten(); t = int(v.sum())
            d.text((400, 384), '(' + '+'.join(str(int(k)) for k in v) + ') / 9', font=FM(19), fill=INK, anchor='mm')
            d.text((400, 416), f'= {t} / 9 = {out[y-1, x-1]}  (média dos 9 pixels)', font=FM(19), fill=AC, anchor='mm'); fr.append(im); du.append(110); n += 1
    du[-1] = 1800; return save(fr, du, 'suavizar_passo.gif')
def g_blur_wipe():
    w, h = 440, 330; pad = np.pad(S, 1, mode='edge'); B = np.round(sum(pad[1 + dy:1 + dy + 72, 1 + dx:1 + dx + 96] for dy in (-1, 0, 1) for dx in (-1, 0, 1)) / 9)
    up = lambda a: np.array(Image.fromarray(np.clip(a, 0, 255).astype('uint8')).resize((w, h), Image.NEAREST)).astype(float)
    W = wipe_panel(up(S), up(B), w, h); fr, du = [], []
    for t in list(np.linspace(0, 1, 14)) + list(np.linspace(1, 0, 14))[1:]:
        im = W(t); d = ImageDraw.Draw(im)
        for tx, xx, an in (('antes', 8, 'lm'), ('depois: média 3×3', w - 8, 'rm')): d.rectangle([xx - (0 if an == 'lm' else 190), h - 30, xx + (80 if an == 'lm' else 0), h - 4], fill=INK); d.text((xx, h - 17), tx, font=F(16, True), fill='white', anchor=an)
        fr.append(im); du.append(110 if 0 < t < 1 else 800)
    return save(fr, du, 'suavizar_foto.gif')

G = dict(zoom=g_zoom(), matriz=g_matrix(), cinza=g_gray(), receita=g_recipe(), regras=g_rules(), usos=g_uses(), passo=g_edge_scan(), foto=g_edge_wipe(), passo2=g_blur_scan(), foto2=g_blur_wipe())

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
rgbc = lambda t: RGBColor(*t)
def tb(s, paras, x, y, w, h, size=22, bold=False, color=INK):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame; t.word_wrap = True
    for i, p in enumerate(paras if isinstance(paras, list) else [paras]):
        para = t.paragraphs[0] if i == 0 else t.add_paragraph(); para.space_after = Pt(10); r = para.add_run(); r.text = p; r.font.size = Pt(size); r.font.bold = bold; r.font.name = 'Calibri'; r.font.color.rgb = rgbc(color)
def slide(title, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgbc(BG)
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.18)); b.fill.solid(); b.fill.fore_color.rgb = rgbc(AC); b.line.fill.background()
    tb(s, title, .5, .35, 12.3, .9, 38, True, AC); s.notes_slide.notes_text_frame.text = notes; return s
def pic(s, path, x, y, mw, mh):
    im = Image.open(path); r = min(mw / im.width, mh / im.height); s.shapes.add_picture(path, Inches(x), Inches(y), Inches(im.width * r), Inches(im.height * r))
HINT = 'Inicie a apresentação (F5) para as animações rodarem. '
s = slide('A imagem como uma grade', HINT + 'Mostre o zoom: a foto é feita de quadradinhos. Pergunte quantos pixels a turma acha que a foto tem. Na plataforma cada aluno sobe a própria foto, que é convertida para P&B.'); tb(s, ['Toda foto digital é uma grade de quadradinhos: os pixels.', 'Quanto mais pixels, mais detalhe.', 'Na plataforma, a sua foto (convertida para preto e branco) vira uma grade de 96 × 72 pixels.'], .5, 1.6, 4.6, 4); tb(s, 'Dica: modo apresentação (F5) para ver a animação.', .5, 6.7, 6, .5, 14, False, MU); pic(s, G['zoom'], 5.3, 1.5, 7.6, 5.6)
s = slide('A matriz de pixels', HINT + 'Cada quadrado guarda um número. A mesma matriz que o aluno vai ver na lupa da plataforma.'); tb(s, ['Cada pixel guarda um número.', '0 = preto', '255 = branco', 'Valores no meio = tons de cinza.'], .5, 1.6, 4.2, 4); pic(s, G['matriz'], 4.9, 1.5, 8.1, 5.6)
s = slide('A escala de cinza', HINT + 'Com 8 bits temos 2^8 = 256 valores possíveis.'); tb(s, ['Do preto ao branco, passamos por muitos tons.', 'Com 8 bits há 256 valores possíveis: de 0 a 255.', 'Foto colorida → cinza: 0,299·R + 0,587·G + 0,114·B (o verde pesa mais, pois o olho enxerga melhor o verde).'], .5, 1.6, 4.2, 4); pic(s, G['cinza'], 4.9, 1.5, 8.1, 5.6)
s = slide('O que é um filtro?', HINT + 'Analogia: uma receita. Entram os pixels, a regra transforma cada um, saem novos pixels. Note o caso 230 + 40: a regra tem um limite.'); tb(s, ['Um filtro é uma regra (uma receita) aplicada a cada pixel.', 'Entrada → regra → saída.', 'Regras também precisam de limites: o branco não passa de 255.'], .5, 1.6, 4.4, 4); pic(s, G['receita'], 5.0, 1.5, 8.0, 5.6)
s = slide('Uma mesma imagem, diferentes regras', HINT + 'Apenas demonstração: na plataforma os alunos farão isso sozinhos.'); tb(s, ['Clarear: p + 40', 'Escurecer: p − 40', 'Inverter: 255 − p', 'Mesma foto, regras diferentes, resultados diferentes.'], .5, 1.6, 4.4, 4); pic(s, G['regras'], 5.2, 1.5, 7.8, 5.6)
s = slide('Onde filtros são úteis?', HINT + 'Exemplos reais. Chame atenção ao trade-off: suavizar remove ruído mas apaga detalhes.'); pic(s, G['usos'], .6, 1.5, 12.1, 4.3); tb(s, 'Atenção: um filtro pode ajudar um objetivo e prejudicar outro (suavizar tira o granulado, mas também apaga detalhes).', .6, 6.0, 12.1, 1, 20, True, AC)
s = slide('Agora é com você', 'Explique que o professor libera cada etapa e que, ao avançar, o tempo da etapa anterior termina.')
for i, (n, t) in enumerate([('Aplicar', 'experimente filtros e publique'), ('Analisar', 'minijogo Filtro Misterioso: descubra o filtro e suba no ranking'), ('Avaliar', 'julgue regras e limites'), ('Criar', 'monte seu próprio filtro')]):
    x = .6 + i * 3.1; c = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2), Inches(2.9), Inches(2.6)); c.fill.solid(); c.fill.fore_color.rgb = RGBColor(255, 255, 255); c.line.color.rgb = rgbc(AC)
    tb(s, f'{i+1}. {n}', x + .15, 2.2, 2.6, .9, 22, True, AC); tb(s, t, x + .15, 3.1, 2.6, 1.4, 20)
tb(s, 'O professor libera cada etapa. Quando ele avança, o tempo da etapa termina — aproveite bem!', .6, 5.2, 12, 1, 22, True)
s = slide('Explore outros filtros: contornar', HINT + 'Extra: detecção de bordas. A cada pixel, compara os vizinhos: onde há grande diferença, há um contorno. Cada filtro usa a vizinhança.'); pic(s, G['passo'], .4, 1.4, 7.6, 4.3); pic(s, G['foto'], 8.2, 1.4, 4.8, 4.3)
tb(s, ['Filtros de contorno comparam cada pixel com seus vizinhos: onde a diferença é grande, há uma borda (as veias das folhas).', 'Na etapa Criar você vai usar os blocos Borda e Suavizar. Desafio: que outros filtros você consegue imaginar? Nitidez, contraste, inverter só as bordas...'], .5, 5.8, 12.3, 1.5, 19)
s = slide('Explore outros filtros: suavizar', HINT + 'Extra: suavização (desfoque). A mesma janela 3×3 do filtro de contorno, mas a conta é outra: a média dos 9 pixels. Diferenças bruscas diminuem: some ruído, mas também detalhes. Na etapa Criar os alunos podem combinar Suavizar e Borda.'); pic(s, G['passo2'], .4, 1.4, 7.6, 4.3); pic(s, G['foto2'], 8.2, 1.4, 4.8, 4.3)
tb(s, ['Suavizar troca cada pixel pela média dele e dos 8 vizinhos: as diferenças bruscas diminuem (o ruído some, mas os detalhes também).', 'Compare com contornar: mesma janela 3×3, outra conta. Na etapa Criar, o que muda se você suavizar antes de achar as bordas?'], .5, 5.8, 12.3, 1.5, 19)
prs.save('slides/aula_pdi.pptx'); print('ok', len(prs.slides._sldIdLst), 'slides')
